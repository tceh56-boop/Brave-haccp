# -*- coding: utf-8 -*-
"""
Сборка дерева папок ТТК кафе: python3 build_ttk_tree.py <сборник.docx> <папка вывода>

Результат в папке вывода:
  ТТК кафе/                         — дерево папок: карты по одной в файле, в каждой папке её СТО
  ТТК кафе.zip                      — то же дерево архивом (загрузить на Google Диск и распаковать)
  Архитектура дерева ТТК и СТО.docx — принципы, дерево, СТО всех папок, отклонения
  Реестр ТТК и СТО.xlsx             — карта → папка → СТО → срок в карте → вывод
  ttk_tree.json                     — дерево и нормы для MonChef (генератор правил сроков)
"""
import copy, io, json, os, re, shutil, sys, zipfile
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from ttk_tree import FOLDERS, NORMS, ORG, ORG_SHORT, YEAR, ROOT, COMMON_DOCS, folder_dir, sto_code, regime_for, check

SRC, OUT = sys.argv[1], sys.argv[2]
TITLE_RE = re.compile(r'^ТЕХНИКО-ТЕХНОЛОГИЧЕСКАЯ КАРТА № (\d+)')


def text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


# ------------------------------------------------------------------ разбор сборника
master = docx.Document(SRC)
body = master.element.body
els = [e for e in body.iterchildren() if e.tag in (qn('w:p'), qn('w:tbl'))]
titles = [i for i, e in enumerate(els) if e.tag == qn('w:p') and TITLE_RE.match(text_of(e).strip())]
starts = []
for i in titles:
    j = i - 1
    while j >= 0 and els[j].tag == qn('w:p') and not text_of(els[j]).strip():
        j -= 1
    starts.append(j if j >= 0 and els[j].tag == qn('w:tbl') and 'УТВЕРЖДАЮ' in text_of(els[j]) else i)
cards = {}
for k, i in enumerate(titles):
    n = int(TITLE_RE.match(text_of(els[i]).strip()).group(1))
    end = starts[k + 1] if k + 1 < len(starts) else len(els)
    rng = els[starts[k]:end]
    name = text_of(els[i + 1]).strip().strip('«»').strip()
    paras = [text_of(e) for e in rng if e.tag == qn('w:p')]
    st, on = [], False
    for p in paras:
        if re.match(r'^5\.\s', p): on = True
        elif re.match(r'^6\.\s', p): on = False
        if on and not re.match(r'^5\.\s', p): st.append(p.strip())
    store = ' '.join(st)
    m = re.search(r'Условия и сроки хранения:\s*(.*)', store)
    cards[n] = {'n': n, 'name': name, 'els': rng, 'store': store, 'storeShort': (m.group(1) if m else store).strip()}
assert len(cards) == 234, 'карт в сборнике: %d' % len(cards)
assign = check()
front = els[:starts[0]]

# раздел сборника (содержание)
section = {}
cur = ''
for e in front:
    if e.tag == qn('w:p') and text_of(e).strip():
        cur = text_of(e).strip()
    elif e.tag == qn('w:tbl'):
        for tr in e.iter(qn('w:tr')):
            m = re.match(r'ТТК № (\d+)', text_of(tr))
            if m: section[int(m.group(1))] = cur


# ------------------------------------------------------------------ сроки в карте
def parse_terms(s):
    """Сроки из п. 5 карты: {'chill': ч или None, 'frozen': сут или None, 'immediate': bool, 'dressed': ч заправленного салата}."""
    out = {'chill': None, 'frozen': None, 'immediate': bool(re.search(r'сразу|хранение не допускается', s, re.I)), 'dressed': None}
    mode = 'chill'
    tok = re.finditer(r'(−\s?18|-\s?18|заморож|\+\s?\d+|(\d+(?:[.,]\d+)?)\s*(?:[–-]\s*(\d+)\s*)?(ч\b|час|сут|дн))', s)
    for t in tok:
        g = t.group(0)
        if re.match(r'[−-]\s?18|заморож', g): mode = 'frozen'; continue
        if g.startswith('+'):
            mode = 'chill'; continue
        v = float((t.group(3) or t.group(2)).replace(',', '.'))
        unit = t.group(4)
        hours = v * 24 if unit.startswith(('сут', 'дн')) else v
        pre = s[max(0, t.start() - 40):t.start()]
        if re.search(r'заправленн', pre) and not re.search(r'без заправки', pre):
            out['dressed'] = hours
            continue
        if mode == 'frozen':
            out['frozen'] = max(out['frozen'] or 0, hours / 24)
        else:
            out['chill'] = max(out['chill'] or 0, hours)
    return out


def fmt_h(h):
    if h is None: return '—'
    if h > 72 and h % 24 == 0: return '%d сут' % (h // 24)
    return ('%g ч' % h).replace('.', ',')


def verdict(f, c):
    r = regime_for(f, c['n'])
    t = parse_terms(c['store'])
    c['terms'] = t
    issues = []
    if f['kind'] == 'dish':
        if not t['immediate']:
            issues.append('в карте не указано «реализуют сразу, хранение не допускается»')
    else:
        if t['chill'] is not None and r['hours'] is not None and t['chill'] > r['hours'] + 0.01:
            issues.append('срок охлаждённого %s больше нормы СТО %s' % (fmt_h(t['chill']), fmt_h(r['hours'])))
        if t['frozen'] is not None:
            if r['frozen'] is None:
                issues.append('заморозка %d сут — СТО для группы заморозку не предусматривает' % t['frozen'])
            elif t['frozen'] > r['frozen'] + 0.01:
                issues.append('срок замороженного %d сут больше нормы СТО %d сут' % (t['frozen'], r['frozen']))
        if t['dressed'] is not None:
            issues.append('хранение заправленного салата %s — по СТО салат заправляют перед подачей' % fmt_h(t['dressed']))
        if t['chill'] is None and t['frozen'] is None and not t['immediate']:
            issues.append('срок хранения в карте не указан')
    if issues:
        todo = 'привести срок в карте к норме СТО либо обосновать больший срок лабораторными исследованиями (МУК 4.2.1847-04) и внести изменение в СТО'
        if any('заправленн' in x for x in issues): todo = 'исправить п. 5 карты: заправлять перед подачей, хранить только незаправленную заготовку'
        if any('не указан' in x for x in issues): todo = 'указать в п. 5 карты условия и срок по СТО папки'
        if f['kind'] == 'dish': todo = 'дописать в п. 5 карты: блюдо собирают по заказу и реализуют сразу; покупные компоненты хранят по условиям изготовителя, собственные п/ф — по СТО их папок'
    else:
        todo = ''
    return r, issues, todo


rows = []
for f in FOLDERS:
    for n in f['cards']:
        c = cards[n]
        r, issues, todo = verdict(f, c)
        rows.append({'n': n, 'name': c['name'], 'section': section.get(n, ''), 'folder': f, 'regime': r, 'issues': issues, 'todo': todo, 'store': c['storeShort'], 'terms': c['terms']})
rows.sort(key=lambda x: x['n'])
dev = [x for x in rows if x['issues']]


# ------------------------------------------------------------------ оформление docx
def base_doc():
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    st.paragraph_format.space_after = Pt(4)
    for s in ('Heading 1', 'Heading 2', 'Heading 3', 'Title'):
        h = d.styles[s]; h.font.name = 'Times New Roman'; h.font.color.rgb = RGBColor(0x5a, 0x10, 0x1e)
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    d.styles['Heading 1'].font.size = Pt(15); d.styles['Heading 2'].font.size = Pt(13); d.styles['Heading 3'].font.size = Pt(12)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(1.5), Cm(2), Cm(2)
    return d


def para(d, text, bold=False, align=None, size=None, italic=False):
    p = d.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    return p


def bullet(d, text):
    p = d.add_paragraph(style='List Bullet'); p.add_run(text); return p


def table(d, head, data, widths=None, size=10):
    t = d.add_table(rows=1, cols=len(head)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(size)
        sh = c._tc.get_or_add_tcPr(); fill = sh.makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:color'): 'auto', qn('w:fill'): 'EFE6E0'}); sh.append(fill)
    for row in data:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''; r = cells[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths): t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t


def regime_rows(f):
    out = []
    for r in f['regimes']:
        if f['kind'] == 'dish':
            out.append([r['what'], r['temp'], r['note']])
        else:
            out.append([r['what'], r['temp'], fmt_h(r['hours']), ('%d сут' % r['frozen']) if r['frozen'] else 'не допускается', r['note'] or '—'])
    return out


# ------------------------------------------------------------------ СТО папки
def sto_doc(f):
    d = base_doc()
    t = d.add_table(rows=1, cols=2); t.style = 'Table Grid'
    t.rows[0].cells[0].text = ORG + '\n(наименование организации)'
    t.rows[0].cells[1].text = 'УТВЕРЖДАЮ\n' + ORG + '\n__________________\n«___» ____________ ' + YEAR + ' г.'
    para(d, '')
    para(d, 'СТАНДАРТ ОРГАНИЗАЦИИ', True, WD_ALIGN_PARAGRAPH.CENTER, 14)
    para(d, sto_code(f), True, WD_ALIGN_PARAGRAPH.CENTER, 14)
    para(d, ('Блюда. ' if f['kind'] == 'dish' else 'Полуфабрикаты. ') + f['name'], True, WD_ALIGN_PARAGRAPH.CENTER, 13)
    para(d, 'Требования к изготовлению, условиям и срокам хранения, реализации и контролю', False, WD_ALIGN_PARAGRAPH.CENTER)
    para(d, 'Папка ТТК: «' + ROOT + '/' + '/'.join(folder_dir(f)) + '»', False, WD_ALIGN_PARAGRAPH.CENTER, 10, True)
    para(d, 'Дата введения: «___» ____________ ' + YEAR + ' г.', False, WD_ALIGN_PARAGRAPH.CENTER, 10)

    d.add_heading('1. Область применения', 2)
    para(d, '1.1. Стандарт распространяется на ' + ('блюда' if f['kind'] == 'dish' else 'полуфабрикаты') + ' группы «' + f['name'] + '», изготовляемые ' + ORG_SHORT +
         ' (' + f['shops'] + '), и обязателен для всех технико-технологических карт, находящихся в папке «' + folder_dir(f)[1] + '». Карта, помещённая в папку, подчиняется этому стандарту; требования карты могут быть строже стандарта, но не мягче.')
    para(d, '1.2. Перечень карт группы на дату утверждения (' + str(len(f['cards'])) + '):')
    table(d, ['№ ТТК', 'Наименование', 'Режим СТО'], [[n, cards[n]['name'], regime_for(f, n)['id'] + ' — ' + regime_for(f, n)['what']] for n in f['cards']], [1.6, 9.5, 5.9], 9)

    d.add_heading('2. Нормативные ссылки', 2)
    for c, n in NORMS: bullet(d, c + ' ' + n + ';')
    para(d, 'Сроки годности продукции собственного изготовления устанавливает изготовитель (ТР ТС 021/2011); СанПиН 2.3/2.4.3590-20 устанавливает требования к условиям изготовления, хранения, раздачи и маркировке. Сроки настоящего стандарта — нормы предприятия; установление более длительного срока допускается только по результатам исследований по МУК 4.2.1847-04.', size=10, italic=True)

    d.add_heading('3. Характеристика группы и опасные факторы', 2)
    para(d, '3.1. Основные биологические опасности группы: ' + f['hazards'])
    para(d, '3.2. Аллергены указывают в ТТК и в меню (информирование потребителя).')

    d.add_heading('4. Требования к сырью', 2)
    para(d, '4.1. Сырьё и пищевые продукты принимают при наличии документов, подтверждающих безопасность и качество (декларация, сертификат, ветеринарные сопроводительные документы для подконтрольной продукции), с маркировкой; входной контроль — с записью в журнале бракеража скоропортящейся продукции.')
    para(d, '4.2. Не допускаются: продукция без маркировки и документов, с истёкшим сроком годности, с признаками порчи; непотрошёная птица; яйца с загрязнённой и повреждённой скорлупой, из хозяйств, неблагополучных по сальмонеллёзам.')
    para(d, '4.3. Полуфабрикаты собственного изготовления используют в пределах срока, установленного их СТО (по маркировке).')

    d.add_heading('5. Требования к изготовлению. Критические контрольные точки', 2)
    para(d, '5.1. Изготовление — по ТТК, с соблюдением поточности (сырое и готовое раздельно), на маркированном инвентаре и оборудовании.')
    for i, x in enumerate(f['tech']): para(d, '5.%d. %s' % (i + 2, x[0].upper() + x[1:]))
    para(d, '5.%d. Критические пределы, процедуры мониторинга и корректирующие действия ККТ — по плану ХАССП предприятия (ТР ТС 021/2011); измерения записывают в журналы системы MonChef.' % (len(f['tech']) + 2))

    d.add_heading('6. Условия и сроки хранения' if f['kind'] == 'pf' else '6. Условия реализации и хранения', 2)
    if f['kind'] == 'pf':
        table(d, ['Режим / вид полуфабриката', 'Условия хранения', 'Срок, охлаждённый', 'Срок, замороженный', 'Примечание'], [[r['id'] + '. ' + x[0]] + x[1:] for r, x in zip(f['regimes'], regime_rows(f))], [4.2, 3.6, 2.3, 2.3, 4.6], 9)
        para(d, '6.1. Срок исчисляется с момента окончания технологического процесса (для охлаждаемых — с момента окончания охлаждения).')
        para(d, '6.2. Хранят в закрытой пищевой таре (гастроёмкости, контейнеры, вакуумные пакеты), раздельно от сырья, на стеллажах не ниже 15 см от пола; холодильное оборудование — с термометрами, температуру записывают в журнал.')
        para(d, '6.3. Размороженный полуфабрикат повторно не замораживают; размораживают в холодильнике при +2…+6 °С.')
        para(d, '6.4. Смешивание партий с разными датами изготовления не допускается; остатки предыдущей партии не добавляют к новой.')
    else:
        table(d, ['Вид', 'Температура при отпуске', 'Реализация'], regime_rows(f), [4, 7, 6], 9)
        para(d, '6.1. Температура при отпуске (раздаче) — по СанПиН 2.3/2.4.3590-20 и ГОСТ 30390-2013: ' + '; '.join([s for s in ['горячие супы, соусы, напитки — не ниже +75 °С', 'вторые блюда и гарниры — не ниже +65 °С', 'холодные супы и напитки — не выше +14 °С']]) + '.')
        para(d, '6.2. Блюда готовят по заказу; оставленные в зале и не реализованные блюда повторно не используются.')

    d.add_heading('7. Маркировка', 2)
    para(d, '7.1. На таре каждого полуфабриката и заготовки — маркировочный ярлык: наименование, дата и время изготовления, срок годности (дата, время), условия хранения, ФИО изготовителя. Ярлык печатает MonChef при выпуске партии (этикетка с QR-кодом партии).' if f['kind'] == 'pf' else
         '7.1. Блюда на вынос упаковывают в одноразовую тару с указанием наименования, даты и времени изготовления, срока и условий реализации, информации об аллергенах.')

    d.add_heading('8. Контроль', 2)
    para(d, '8.1. Органолептическая оценка (бракераж) готовой продукции — перед отпуском, с записью в журнале бракеража готовой пищевой продукции (ГОСТ 31986-2012).')
    para(d, '8.2. Лабораторный контроль микробиологических показателей (ТР ТС 021/2011, приложение 1 — продукция общественного питания соответствующей группы) — по программе производственного контроля (ППК).')
    para(d, '8.3. Ежедневно — контроль сроков по маркировке: продукция с истёкшим сроком списывается актом и не используется. MonChef ставит задачи «Использовать в первую очередь» и «Списать просроченную партию».')

    if f['kind'] == 'pf' or any(x['issues'] for x in rows if x['folder'] is f):
        d.add_heading('9. Соответствие карт группы стандарту', 2)
        fr = [x for x in rows if x['folder'] is f]
        table(d, ['№ ТТК', 'Наименование', 'Срок в карте (п. 5)', 'Норма СТО', 'Вывод'],
              [[x['n'], x['name'], x['store'][:180], (fmt_h(x['regime']['hours']) + ((' / зам. %d сут' % x['regime']['frozen']) if x['regime']['frozen'] else '')) if f['kind'] == 'pf' else 'сразу',
                'соответствует' if not x['issues'] else '; '.join(x['issues']) + '. Что сделать: ' + x['todo']] for x in fr], [1.4, 3.8, 5.2, 2.3, 4.3], 8)

    d.add_heading('10. Порядок изменения', 2)
    para(d, '10.1. Новая карта группы помещается в папку и подчиняется стандарту с даты помещения. Изменение сроков и условий стандарта — приказом руководителя; увеличение срока — по результатам исследований (МУК 4.2.1847-04), протоколы прилагаются к стандарту.')
    para(d, '10.2. Стандарт пересматривается при изменении нормативных документов, ассортимента, оборудования и не реже одного раза в год при актуализации плана ХАССП.')
    para(d, '')
    para(d, 'Разработал: ____________ / ____________________ /        Согласовано (ответственный за ХАССП): ____________ / ____________________ /', size=10)
    return d


# ------------------------------------------------------------------ вывод
tree = os.path.join(OUT, ROOT)
if os.path.exists(tree): shutil.rmtree(tree)
os.makedirs(os.path.join(tree, COMMON_DOCS, 'СТО'))
os.makedirs(os.path.join(tree, '99 Архив (выведенные из меню и прежние версии)'))

# шаблон карты: пакет сборника без содержимого (стили, колонтитулы, поля сохраняются)
tpl = docx.Document(SRC)
for e in list(tpl.element.body.iterchildren()):
    if e.tag != qn('w:sectPr'): tpl.element.body.remove(e)
tbuf = io.BytesIO(); tpl.save(tbuf); TPL = tbuf.getvalue()


def strip_breaks(el_list):
    for e in el_list[:2] + el_list[-2:]:
        for br in e.iter(qn('w:br')):
            if br.get(qn('w:type')) == 'page': br.getparent().remove(br)
        for pb in e.iter(qn('w:pageBreakBefore')): pb.getparent().remove(pb)


def safe(s):
    return re.sub(r'\s+', ' ', re.sub(r'[\\/:*?"<>|«»]', '', s)).strip()[:90]


for f in FOLDERS:
    top, sub = folder_dir(f)
    fd = os.path.join(tree, top, sub)
    os.makedirs(fd, exist_ok=True)
    for n in f['cards']:
        d = docx.Document(io.BytesIO(TPL)); b = d.element.body; sp = b.find(qn('w:sectPr'))
        cl = [copy.deepcopy(e) for e in cards[n]['els']]
        strip_breaks(cl)
        for e in cl: sp.addprevious(e)
        d.save(os.path.join(fd, 'ТТК %03d %s.docx' % (n, safe(cards[n]['name']))))
    sd = sto_doc(f)
    fn = '%s %s.docx' % (sto_code(f), safe(f['name']))
    sd.save(os.path.join(fd, '_' + fn))
    sd.save(os.path.join(tree, COMMON_DOCS, 'СТО', fn))
shutil.copy(SRC, os.path.join(tree, COMMON_DOCS, 'Сборник ТТК № 1–234 (исходный).docx'))

# реестр xlsx
wb = openpyxl.Workbook()
hdr_fill, wrap = PatternFill('solid', fgColor='EFE6E0'), Alignment(wrap_text=True, vertical='top')


def sheet(ws, head, data, widths):
    ws.append(head)
    for c in ws[1]: c.font = Font(bold=True); c.fill = hdr_fill; c.alignment = wrap
    for r in data: ws.append(r)
    for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = wrap
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions


ws = wb.active; ws.title = 'Реестр ТТК'
sheet(ws, ['№ ТТК', 'Наименование', 'Раздел сборника', 'Папка', 'Путь в дереве', 'СТО', 'Режим СТО', 'Норма СТО: охлаждённый', 'Норма СТО: замороженный',
           'Срок в карте (п. 5)', 'В карте: охлаждённый, ч', 'В карте: замороженный, сут', 'Вывод', 'Что сделать'],
      [[x['n'], x['name'], x['section'], x['folder']['code'] + ' ' + x['folder']['name'], ROOT + '/' + '/'.join(folder_dir(x['folder'])), sto_code(x['folder']),
        x['regime']['id'] + ' — ' + x['regime']['what'], 'сразу' if x['folder']['kind'] == 'dish' else fmt_h(x['regime']['hours']),
        ('%d сут' % x['regime']['frozen']) if x['regime']['frozen'] else '—', x['store'], x['terms']['chill'] if x['terms']['chill'] is not None else '',
        x['terms']['frozen'] if x['terms']['frozen'] is not None else '', 'соответствует' if not x['issues'] else '; '.join(x['issues']), x['todo']] for x in rows],
      [7, 34, 26, 30, 50, 13, 34, 12, 12, 60, 10, 10, 50, 50])
red = PatternFill('solid', fgColor='F8D7D3')
for row in ws.iter_rows(min_row=2):
    if row[12].value != 'соответствует':
        for c in row: c.fill = red
ws2 = wb.create_sheet('Папки и СТО')
sheet(ws2, ['Код', 'Папка', 'Путь', 'Вид', 'Цеха', 'Карт', 'СТО', 'Режимы хранения / реализации', 'Отклонений'],
      [[f['code'], f['name'], ROOT + '/' + '/'.join(folder_dir(f)), 'блюда' if f['kind'] == 'dish' else 'полуфабрикаты', f['shops'], len(f['cards']), sto_code(f),
        '\n'.join(r['id'] + '. ' + r['what'] + ': ' + r['temp'] + ('' if f['kind'] == 'dish' else '; охл. ' + fmt_h(r['hours']) + '; зам. ' + (('%d сут' % r['frozen']) if r['frozen'] else 'не допускается')) + (('; ' + r['note']) if r['note'] else '') for r in f['regimes']),
        sum(1 for x in rows if x['folder'] is f and x['issues'])] for f in FOLDERS],
      [7, 40, 60, 14, 26, 7, 14, 90, 11])
ws3 = wb.create_sheet('Отклонения')
sheet(ws3, ['№ ТТК', 'Наименование', 'Папка', 'Срок в карте', 'Отклонение', 'Что сделать'],
      [[x['n'], x['name'], x['folder']['code'], x['store'], '; '.join(x['issues']), x['todo']] for x in dev], [7, 34, 8, 60, 60, 60])
wb.save(os.path.join(OUT, 'Реестр ТТК и СТО.xlsx'))
shutil.copy(os.path.join(OUT, 'Реестр ТТК и СТО.xlsx'), os.path.join(tree, COMMON_DOCS, 'Реестр ТТК и СТО.xlsx'))

# архитектура
d = base_doc()
para(d, ORG, False, WD_ALIGN_PARAGRAPH.CENTER, 11)
para(d, 'АРХИТЕКТУРА ДЕРЕВА ПАПОК ТЕХНИКО-ТЕХНОЛОГИЧЕСКИХ КАРТ', True, WD_ALIGN_PARAGRAPH.CENTER, 15)
para(d, 'и стандарты организации (СТО) групп продукции', True, WD_ALIGN_PARAGRAPH.CENTER, 13)
para(d, 'Сборник ТТК № 1–234 · %d папок · %d СТО · %d карт' % (len(FOLDERS), len(FOLDERS), len(rows)), False, WD_ALIGN_PARAGRAPH.CENTER, 11, True)
d.add_heading('1. Принципы', 1)
for s in ['Папка — это группа продукции с одинаковыми санитарными требованиями к изготовлению, хранению и реализации. Поэтому у каждой папки свой СТО, и любая карта в папке подчиняется ему: при добавлении новой карты условия и сроки её хранения задаются местом в дереве, а не памятью технолога.',
          'Два корня по главному санитарному признаку: «01 Блюда» — готовятся по заказу и не хранятся (СТО задаёт температуру и время реализации); «02 Полуфабрикаты» — хранятся (СТО задаёт режим: температура и предельный срок, охлаждённый и замороженный). Внутри — по виду сырья и степени риска: сырое мясо и рыба отдельно от прошедших тепловую обработку, яичные и кремовые массы — отдельными папками с самыми короткими сроками.',
          'Цех — реквизит карты (в реестре и в MonChef), а не уровень дерева: один полуфабрикат может делать любой цех, а требования к хранению от цеха не зависят.',
          'Глубина дерева — два уровня. Код папки (01.03, 02.06) — её постоянный номер; он же номер СТО (СТО 02.06-2026). Новые группы получают следующий свободный код, коды не переиспользуются.',
          'Внутри папки: карты по одной в файле «ТТК 001 Наименование.docx» (номер — сквозной номер сборника, три цифры для сортировки) и файл СТО папки, начинающийся с «_» — он всегда первый в списке.',
          '«00 Общие документы» — все СТО, реестр карт и исходный сборник; «99 Архив» — выведенные из меню карты и прежние версии (карту не удаляют, а переносят в архив с датой).']:
    bullet(d, s)
d.add_heading('2. Дерево', 1)
lines = [ROOT + '/', '├── ' + COMMON_DOCS + '/  (СТО всех папок, реестр, исходный сборник)']
groups = [('01 Блюда', [f for f in FOLDERS if f['kind'] == 'dish']), ('02 Полуфабрикаты', [f for f in FOLDERS if f['kind'] == 'pf'])]
for gi, (g, fs) in enumerate(groups):
    lines.append(('├── ' if gi == 0 else '└── ') + g + '/')
    for i, f in enumerate(fs):
        nm = f['name'] if len(f['name']) <= 52 else f['name'][:50].rsplit(' ', 1)[0] + '…'
        lines.append(('│   ' if gi == 0 else '    ') + ('└── ' if i == len(fs) - 1 else '├── ') + f['code'] + ' ' + nm + '  [' + str(len(f['cards'])) + ']')
lines.insert(2, '├── 99 Архив (выведенные из меню и прежние версии)/')
for l in lines:
    p = d.add_paragraph(); r = p.add_run(l); r.font.name = 'Consolas'; r.font.size = Pt(8); p.paragraph_format.space_after = Pt(0)
para(d, 'В квадратных скобках — число карт. Номер СТО папки совпадает с её кодом: папка 02.06 — СТО 02.06-' + YEAR + '.', size=9, italic=True)
d.add_heading('3. СТО папок: условия и сроки', 1)
para(d, 'Нормы — предприятия (см. раздел 5). Полный текст каждого СТО — в папке группы (файл «_СТО …») и в «00 Общие документы/СТО».', size=10, italic=True)
for g, fs in groups:
    d.add_heading(g, 2)
    if g.startswith('01'):
        table(d, ['Папка', 'СТО', 'Карт', 'Температура при отпуске', 'Реализация'],
              [[f['code'] + ' ' + f['name'], sto_code(f), len(f['cards']), f['regimes'][0]['temp'], f['regimes'][0]['note']] for f in fs], [4.5, 2.3, 1.2, 5, 4], 8)
    else:
        data = []
        for f in fs:
            for r in f['regimes']:
                data.append([f['code'] + ' ' + f['name'] if r is f['regimes'][0] else '', r['id'] + '. ' + r['what'], r['temp'], fmt_h(r['hours']), ('%d сут' % r['frozen']) if r['frozen'] else '—'])
        table(d, ['Папка', 'Режим', 'Условия', 'Охлаждённый', 'Замороженный'], data, [4.2, 5.2, 3.8, 2, 2], 8)
d.add_heading('4. Карты по папкам и отклонения от СТО', 1)
para(d, 'Карт соответствует своему СТО: %d из %d. Отклонений: %d — их нужно исправить в карте или обосновать (раздел 5).' % (len(rows) - len(dev), len(rows), len(dev)))
for g, fs in groups:
    d.add_heading(g, 2)
    for f in fs:
        para(d, f['code'] + ' ' + f['name'] + ' — ' + sto_code(f), True, size=11)
        para(d, ', '.join('№ %d %s' % (n, cards[n]['name']) for n in f['cards']), size=9)
if dev:
    d.add_heading('Отклонения от СТО', 2)
    table(d, ['№', 'Карта', 'Папка', 'Отклонение', 'Что сделать'], [[x['n'], x['name'], x['folder']['code'], '; '.join(x['issues']), x['todo']] for x in dev], [1, 4, 1.4, 5.6, 5], 8)
d.add_heading('5. Нормативная основа сроков', 1)
for s in ['СанПиН 2.3/2.4.3590-20 устанавливает требования к условиям изготовления, хранения, раздачи и маркировке продукции общественного питания, но не устанавливает перечня сроков годности для продукции собственного изготовления. Срок годности и условия хранения устанавливает изготовитель (ТР ТС 021/2011) — в ТТК и в СТО.',
          'Сроки СТО выбраны не выше обычных для групп продукции значений и совпадают со сроками большинства карт сборника. Отменённый СанПиН 2.3.2.1324-03 (утратил силу с 01.03.2021) в качестве основания не используется.',
          'Срок больше нормы СТО (например, 14 суток для соусов в вакууме, 5–7 суток для су-вида и пастрами) допускается только после исследований по МУК 4.2.1847-04 в аккредитованной лаборатории; по их результатам режим вносится в СТО приказом.',
          'Требования о температуре при отпуске — СанПиН 2.3/2.4.3590-20 и ГОСТ 30390-2013; об обеззараживании рыбы, употребляемой без тепловой обработки, — СанПиН 3.3686-21 (заменил СанПиН 3.2.3215-14, на который ссылаются карты № 49, 117, 139 — ссылку в картах обновить).']:
    bullet(d, s)
d.add_heading('6. Нормативные документы', 2)
for c, n in NORMS: bullet(d, c + ' ' + n)
d.add_heading('7. Работа с деревом в MonChef', 1)
for s in ['Дерево загружается в папку ТТК на Google Диске (архив «ТТК кафе.zip» — распаковать в корень папки ТТК).',
          'Экран MonChef «Меню и ТТК» → «Дерево папок и СТО»: создаёт недостающие папки на Диске и загружает СТО как правила сроков (по шаблону пути папки и, для особых режимов, по наименованию полуфабриката). Правила загружаются черновиками: в действие их вводит ответственный за ХАССП после утверждения СТО — тогда этикетки партий полуфабрикатов получают срок годности из СТО их папки.',
          'Новая карта: создать в нужной папке → срок и условия берутся из СТО → в п. 5 карты указываются они же или более строгие.']:
    bullet(d, s)
d.save(os.path.join(OUT, 'Архитектура дерева ТТК и СТО.docx'))
shutil.copy(os.path.join(OUT, 'Архитектура дерева ТТК и СТО.docx'), os.path.join(tree, COMMON_DOCS, 'Архитектура дерева ТТК и СТО.docx'))

# json для MonChef
js = {'root': ROOT, 'year': YEAR, 'common': COMMON_DOCS, 'archive': '99 Архив (выведенные из меню и прежние версии)',
      'folders': [{'code': f['code'], 'name': f['name'], 'kind': f['kind'], 'top': folder_dir(f)[0], 'dir': folder_dir(f)[1], 'sto': sto_code(f), 'shops': f['shops'],
                   'regimes': f['regimes'], 'cards': [{'n': n, 'name': cards[n]['name'], 'regime': regime_for(f, n)['id']} for n in f['cards']]} for f in FOLDERS],
      'deviations': [{'n': x['n'], 'name': x['name'], 'folder': x['folder']['code'], 'issues': x['issues']} for x in dev]}
json.dump(js, open(os.path.join(OUT, 'ttk_tree.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# архив дерева
zp = os.path.join(OUT, ROOT + '.zip')
with zipfile.ZipFile(zp, 'w', zipfile.ZIP_DEFLATED) as z:
    for dp, dn, fn in os.walk(tree):
        if not fn and not dn: z.writestr(os.path.relpath(dp, OUT) + '/', '')
        for x in sorted(fn): z.write(os.path.join(dp, x), os.path.relpath(os.path.join(dp, x), OUT))
print('Карт:', len(rows), '· папок:', len(FOLDERS), '· отклонений:', len(dev))
for x in dev: print('  №%d %s [%s]: %s' % (x['n'], x['name'][:40], x['folder']['code'], '; '.join(x['issues'])))
